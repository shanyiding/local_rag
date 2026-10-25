"""Loader tests: build real Office/HTML files on the fly and read them back."""
from pathlib import Path

import pytest

from localrag.config import Config
from localrag.loaders import load_file
from localrag.pipeline import discover_files, ingest, retrieve
from localrag.store import VectorStore

from .test_pipeline import FakeClient


@pytest.fixture
def docx_file(tmp_path):
    from docx import Document

    d = Document()
    d.add_heading("Refund Policy", level=1)
    d.add_paragraph("Customers may request a refund within thirty days of purchase.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Plan", "Refund window"
    t.cell(1, 0).text, t.cell(1, 1).text = "Premium", "sixty days"
    p = tmp_path / "policy.docx"
    d.save(p)
    return p


@pytest.fixture
def pptx_file(tmp_path):
    from pptx import Presentation

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Roadmap"
    s1.placeholders[1].text = "Launch the beta in March"
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Budget"
    s2.placeholders[1].text = "Marketing gets forty percent"
    s2.notes_slide.notes_text_frame.text = "Mention the hiring freeze here"
    p = tmp_path / "deck.pptx"
    prs.save(p)
    return p


@pytest.fixture
def xlsx_file(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Region", "Quarter", "Revenue"])
    ws.append(["Europe", "Q1", 1200])
    ws.append(["Asia", "Q2", 3400.5])
    wb.create_sheet("Empty")
    p = tmp_path / "sales.xlsx"
    wb.save(p)
    return p


@pytest.fixture
def html_file(tmp_path):
    p = tmp_path / "page.html"
    p.write_text(
        "<html><head><title>Zebra Facts</title><style>.x{color:red}</style></head>"
        "<body><script>var secret=1;</script><h1>Stripes</h1>"
        "<p>Zebras have black and white stripes.</p>"
        "<table><tr><td>Speed</td><td>65 km/h</td></tr></table></body></html>"
    )
    return p


def test_docx_reads_paragraphs_and_tables(docx_file):
    (sec,) = load_file(docx_file)
    assert "Refund Policy" in sec.text
    assert "thirty days" in sec.text
    assert "Premium | sixty days" in sec.text


def test_pptx_gives_one_section_per_slide_with_notes(pptx_file):
    secs = load_file(pptx_file)
    assert [s.locator for s in secs] == ["slide 1", "slide 2"]
    assert "Launch the beta" in secs[0].text
    assert "hiring freeze" in secs[1].text
    assert secs[1].page == 2


def test_xlsx_turns_rows_into_labelled_pairs_and_skips_empty_sheets(xlsx_file):
    secs = load_file(xlsx_file)
    assert len(secs) == 1 and secs[0].locator == "sheet Sales"
    assert "Region: Europe; Quarter: Q1; Revenue: 1200" in secs[0].text
    assert "Revenue: 3400.5" in secs[0].text


def test_html_strips_scripts_and_styles_keeps_title(html_file):
    (sec,) = load_file(html_file)
    assert "Zebra Facts" in sec.text and "black and white stripes" in sec.text
    assert "secret" not in sec.text and "color:red" not in sec.text
    assert "Speed" in sec.text and "65 km/h" in sec.text


def test_end_to_end_ingest_and_cite_slide(tmp_path, pptx_file, docx_file, xlsx_file, html_file):
    docs = tmp_path / "docs"
    docs.mkdir()
    for f in (pptx_file, docx_file, xlsx_file, html_file):
        (docs / f.name).write_bytes(f.read_bytes())
    (docs / "~$lock.docx").write_bytes(b"junk")      # Office lock file: must be ignored
    (docs / "old.doc").write_bytes(b"legacy")         # legacy: must warn, not crash

    cfg = Config(docs_dir=str(docs), db_dir=str(tmp_path / "db"), query_prefix="",
                 doc_prefix="", min_similarity=0.1)
    assert "~$lock.docx" not in [p.name for p in discover_files(cfg)]

    logs: list[str] = []
    client, store = FakeClient(), VectorStore(cfg.db_dir, "test")
    stats = ingest(cfg, client, store, log=logs.append)
    assert stats.added == 4
    assert any("old.doc" in line and ".docx" in line for line in logs)

    hits = retrieve(cfg, client, store, "hiring freeze budget marketing")
    assert hits[0].source == "deck.pptx"
    assert hits[0].label == "deck.pptx (slide 2)"
    hits = retrieve(cfg, client, store, "Revenue Europe Q1")
    assert hits[0].label == "sales.xlsx (sheet Sales)"
