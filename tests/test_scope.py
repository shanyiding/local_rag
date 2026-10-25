"""Scope parsing, page lookup, hybrid search and automatic re-indexing."""
from pathlib import Path

import pytest

from localrag.config import Config
from localrag.pipeline import INDEX_FORMAT, answer, ingest, retrieve
from localrag.scope import parse_scope
from localrag.store import VectorStore

from .test_pipeline import FakeClient

SOURCES = [
    "eco/ECO320/1/Coase - Social Cost.pdf",
    "eco/ECO320/2/Dewees - SO2 in Sudbury.pdf",
    "eco/ECO320/Midterm/ECO 320H5F - Midterm Test.pdf",
    "sta/sta302/Syllabus.pdf",
    "sta/sta302/Midterms/sta302Fall22-test1.pdf",
]


def test_course_code_and_page_are_detected():
    s = parse_scope("summarize slide 3 of eco320", SOURCES)
    assert s.page == 3
    assert s.sources and all("ECO320" in x or "320" in x for x in s.sources)
    assert not any("sta302" in x for x in s.sources)


def test_spaced_course_code_is_joined():
    s = parse_scope("what is covered in ECO 320", SOURCES)
    assert s.sources and all(x.startswith("eco/") for x in s.sources)


def test_two_courses_fall_back_to_union():
    s = parse_scope("compare eco320 and sta302", SOURCES)
    assert any(x.startswith("eco/") for x in s.sources) and any(x.startswith("sta/") for x in s.sources)


def test_full_file_name_scopes_to_that_file():
    s = parse_scope("slide 4 of Dewees - SO2 in Sudbury", SOURCES)
    assert s.sources == ["eco/ECO320/2/Dewees - SO2 in Sudbury.pdf"] and s.page == 4


def test_plain_words_do_not_scope():
    s = parse_scope("what percent is the midterms worth", SOURCES)
    assert s.sources is None and s.page is None


def test_slide_number_alone_is_not_a_file_code():
    assert parse_scope("explain slide 3", SOURCES).sources is None


# ---------------------------------------------------------------- end to end
def _make_deck(path: Path, texts: list[str]):
    from pptx import Presentation

    prs = Presentation()
    for t in texts:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = t.split(".")[0]
        s.placeholders[1].text = t
    prs.save(path)


@pytest.fixture
def env(tmp_path):
    docs = tmp_path / "docs"
    (docs / "eco320").mkdir(parents=True)
    (docs / "sta302").mkdir()
    _make_deck(docs / "eco320" / "pigou.pptx", ["Intro to externalities.", "Pigouvian taxes explained.", "Coase theorem basics."])
    (docs / "sta302" / "syllabus.txt").write_text(
        "Grading scheme. The midterm is worth 30 percent and the final exam is worth 40 percent."
    )
    (docs / "sta302" / "lecture1.txt").write_text("Regression lines minimise squared residuals across observations.")
    cfg = Config(docs_dir=str(docs), db_dir=str(tmp_path / "db"), query_prefix="", doc_prefix="",
                 min_similarity=0.05, condense_followups=False)
    client, store = FakeClient(), VectorStore(cfg.db_dir, "test")
    ingest(cfg, client, store, log=lambda _: None)
    return cfg, client, store


def test_page_lookup_returns_exactly_that_slide(env):
    cfg, client, store = env
    ans = answer(cfg, client, store, "summarize slide 2 of eco320")
    assert [h.label for h in ans.hits] == ["eco320/pigou.pptx (slide 2)"]
    assert "Pigouvian" in ans.hits[0].text


def test_page_lookup_with_missing_page_says_so(env):
    cfg, client, store = env
    ans = answer(cfg, client, store, "summarize slide 99 of eco320")
    assert ans.refused and "no page or slide 99" in ans.text()


def test_page_lookup_across_many_files_asks_which_one(env):
    cfg, client, store = env
    for n in range(4):
        _make_deck(Path(cfg.docs_dir) / "eco320" / f"deck{n}.pptx", ["a b c.", "second slide."])
    ingest(cfg, client, store, log=lambda _: None)
    ans = answer(cfg, client, store, "summarize slide 2 of eco320")
    text = ans.text()
    assert ans.refused and "Tell me which one" in text and "deck0.pptx" in text


def test_scope_restricts_retrieval_to_the_course(env):
    cfg, client, store = env
    hits = retrieve(cfg, client, store, "midterm worth sta302")
    assert hits and all(h.source.startswith("sta302/") for h in hits)
    assert hits[0].source == "sta302/syllabus.txt"


def test_hybrid_finds_exact_term_that_embeddings_miss(env):
    cfg, client, store = env

    class BlindEmbedder(FakeClient):  # embeddings carry no information at all
        def embed(self, model, texts):
            return [[1.0] + [0.0] * (self.DIM - 1) for _ in texts]

    cfg.min_similarity = 0.0
    hits = retrieve(cfg, BlindEmbedder(), store, "residuals", k=1)
    assert hits[0].source == "sta302/lecture1.txt"


def test_old_index_format_is_reindexed_automatically(env):
    cfg, client, store = env
    # simulate an index built by an older version (format 1)
    res = store.col.get(include=["metadatas"])
    for i in res["ids"]:
        store.col.update(ids=[i], metadatas=[{"fmt": 1}])
    stats = ingest(cfg, client, store, log=lambda _: None)
    assert stats.updated == 3 and stats.skipped == 0
    again = ingest(cfg, client, store, log=lambda _: None)
    assert again.skipped == 3
    assert store.file_info("sta302/syllabus.txt")[1] == INDEX_FORMAT
